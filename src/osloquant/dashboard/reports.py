"""Quarterly reports of research-book companies: found on NewsWeb, read from the report PDF, compared with the previous
quarter, the published consensus and the case's own assumptions, with every kill criterion marked triggered, nearer,
further or unchanged.

No text is generated: every number is read from a table in the report (with its page), from the consensus release,
or from the case model, or derived from those with the formula shown (LTM = last full year + year to date - the same
period a year earlier). Each figure from a key-figure table is checked against the change the report itself prints.

What to read is configured per company in ``~/osloquant-home/reports.yaml`` (private: it holds the kill thresholds):

    SATS:
      key: {revenue: "Total revenues", ebitda_pre16: "EBITDA before impact of IFRS 16", ...}   # quarter / year-ago / % / YTD / YTD-ago / %
      pl: {revenue: "Revenue", personnel: "Personnel expenses"}                              # columns from the table header
      consensus: {mfn: sats}                                                                 # Modular Finance via MFN
      case: {year_row: 3, sheet: Model, revenue: rev, personnel: pers, clubs: loc, ...}       # model rows (model_map rows)
      kill:
        - {criterion: 0, metric: personnel_pct_ltm, op: ">", value: 0.40}   # an example, not a case
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import yaml

from osloquant import judgements as jg
from osloquant.dashboard import feeds as fd

REPORT_CATEGORIES = {1001, 1002}
MFN = "https://mfn.se/a/modularfinance-estimates/{slug}-consensus-estimates-q{q}-{year}"
_PCT = re.compile(r"(-?\d+(?:[.,]\d+)?)\s?%|n\.a\.?", re.I)
_PERIOD = re.compile(r"\b(Q[1-4]|H[12]|YTD|FY)\s?(\d{4})\b|\b(\d{4})\b")


def spec_path() -> Path:
    return jg.home() / "reports.yaml"


def specs() -> dict:
    p = spec_path()
    return (yaml.safe_load(p.read_text(encoding="utf-8")) or {}) if p.exists() else {}


def report_dir(code: str) -> Path:
    d = jg.home() / "reports" / code
    d.mkdir(parents=True, exist_ok=True)
    return d


# ------------------------------------------------------------------------------------------------ numbers in text
def _tokens(s: str) -> list[str]:
    return re.findall(r"-?\d+(?:\.\d+)?", s)


def partitions(tokens: list[str], k: int) -> list[list[float]]:
    """Every way to read the tokens as k numbers written with a space as thousands separator: a number is a token of
    1-3 digits (with an optional minus) followed by tokens of exactly three digits."""
    out: list[list[float]] = []

    def rec(i: int, acc: list[str]):
        if len(acc) == k:
            if i == len(tokens):
                out.append([float(x) for x in acc])
            return
        if i >= len(tokens):
            return
        head = tokens[i]
        digits = head.lstrip("-")
        if "." in head or len(digits) > 3:          # a decimal, or a number written without separators
            rec(i + 1, acc + [head])
            return
        if len(digits) > 1 and digits.startswith("0"):
            return                                   # "085" continues a number; it cannot start one
        num = head
        rec(i + 1, acc + [num])
        j = i + 1
        while j < len(tokens) and re.fullmatch(r"\d{3}", tokens[j]):
            num = num + tokens[j]
            rec(j + 1, acc + [num])
            j += 1

    rec(0, [])
    return out


def _pct_value(m: re.Match) -> float | None:
    return None if m.group(1) is None else float(m.group(1).replace(",", ".")) / 100


@dataclass
class Figure:
    value: float
    page: int
    line: str
    columns: dict = field(default_factory=dict)
    checked: str = ""


def read_key_row(line: str, label: str) -> dict | None:
    """A key-figure row: label, quarter, year-ago quarter, change %, YTD, YTD a year ago, change %. The two numbers in
    each block are chosen so the printed change matches (within 1.5 points)."""
    rest = line[len(label):] if line.startswith(label) else re.sub(rf"^{label}", "", line)
    pcts = list(_PCT.finditer(rest))
    if len(pcts) < 1:
        return None
    blocks, start = [], 0
    for m in pcts[:2]:
        blocks.append((rest[start:m.start()], _pct_value(m)))
        start = m.end()
    out = {}
    for name, (text, printed) in zip(("q", "ytd"), blocks, strict=False):
        cands = partitions(_tokens(text), 2)
        if printed is not None:
            ok = [c for c in cands if c[1] and abs((c[0] / c[1] - 1) - printed) <= 0.015]
            cands = ok or cands
        if len(cands) != 1:
            return None if name == "q" else out
        out[name], out[f"{name}_ly"], out[f"{name}_chg_printed"] = cands[0][0], cands[0][1], printed
    return out


def read_pl_row(line: str, label: str, columns: list[str]) -> dict | None:
    """An income-statement row: label, optional note references, one number per column of the table header."""
    rest = line[len(label):]
    k = len(columns)
    for drop in (r"", r"^\s*\d{1,2}(?:\s*,\s*\d{1,2})*\s+"):
        text = re.sub(drop, "", rest, count=1) if drop else rest
        cands = partitions(_tokens(text), k)
        if len(cands) == 1:
            return dict(zip(columns, cands[0], strict=True))
    return None


def _norm_period(t: str) -> str:
    return re.sub(r"\s+", " ", t.strip())


def pl_columns(header: str) -> list[str]:
    return [_norm_period(m.group(0)) for m in re.finditer(r"(?:Q[1-4]|H[12]|YTD)\s?\d{4}|\b\d{4}\b", header)]


def pdf_lines(path: Path) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    out = []
    for i, page in enumerate(PdfReader(path).pages):
        for line in (page.extract_text() or "").splitlines():
            out.append((i + 1, re.sub(r"\s+", " ", line).strip()))
    return out


def extract(path: Path, spec: dict) -> dict[str, Figure]:
    """Read the configured rows. Key-figure rows: the first matching line in the report (the group table comes before
    the segments). Income-statement rows: the line under the first header that names the periods."""
    lines = pdf_lines(path)
    figs: dict[str, Figure] = {}
    for key, label in (spec.get("key") or {}).items():
        for page, line in lines:
            if re.match(rf"{label}\s+-?\d", line):
                row = read_key_row(line, re.match(label, line).group(0))
                if row and "q" in row:
                    chg = row["q"] / row["q_ly"] - 1 if row["q_ly"] else None
                    checked = ("matches the printed change" if row["q_chg_printed"] is not None and chg is not None
                               and abs(chg - row["q_chg_printed"]) <= 0.015 else "no printed change to check against")
                    figs[key] = Figure(row["q"], page, line, row, checked)
                    break
    header_cols, header_page = None, None
    for page, line in lines:
        cols = pl_columns(line)
        if line.lower().startswith(("notes ", "note ")) and len(cols) >= 3:
            header_cols, header_page = cols, page
            break
    if header_cols:
        for key, label in (spec.get("pl") or {}).items():
            for page, line in lines:
                if page >= header_page and re.match(rf"{label}\s+(\d{{1,2}}(,\s*\d{{1,2}})*\s+)?-?\d", line):
                    row = read_pl_row(line, re.match(label, line).group(0), header_cols)
                    if row:
                        figs[f"pl_{key}"] = Figure(abs(list(row.values())[0]), page, line, {k: abs(v) for k, v in row.items()},
                                                   "read against the table header: " + ", ".join(header_cols))
                        break
    return figs


# ------------------------------------------------------------------------------------------------ derived figures
def period_of(title: str) -> tuple[int, int] | None:
    m = re.search(r"\bQ([1-4])\s?(\d{4})\b", title)
    return (int(m.group(2)), int(m.group(1))) if m else None


def ltm(cols: dict, year: int, q: int) -> float | None:
    """Last twelve months from an income-statement row: FY(y-1) + YTD(y) - YTD(y-1); Q4: the full year; Q1: FY + Q - Q a
    year ago."""
    fy_prev = cols.get(str(year - 1))
    if q == 4:
        return cols.get(str(year))
    if q == 1:
        cur, prev = cols.get(f"Q1 {year}"), cols.get(f"Q1 {year - 1}")
    else:
        cur, prev = cols.get(f"YTD {year}"), cols.get(f"YTD {year - 1}")
    if None in (fy_prev, cur, prev):
        return None
    return fy_prev + cur - prev


FLOWS = ("revenue", "ebitda", "ebitda_pre16", "ebit_pre16", "ebit")


def derive(figs: dict[str, Figure], year: int, q: int, last_fy: dict | None = None) -> dict[str, dict]:
    """The metrics the comparisons and kill rules use, each with its formula. ``last_fy``: the key figures of the
    previous year's Q4 report (their year-to-date column is the full year), for last-twelve-month key figures."""
    m: dict[str, dict] = {}
    for k, f in figs.items():
        if not k.startswith("pl_") and q == 1 and f.columns.get("ytd") is None:
            f.columns.update(ytd=f.columns.get("q"), ytd_ly=f.columns.get("q_ly"))    # Q1: the year to date is the quarter
        if not k.startswith("pl_"):
            m[k] = {"value": f.value, "how": f"report p. {f.page}", "year_ago": f.columns.get("q_ly")}
            if f.columns.get("q_ly"):
                m[f"{k}_yoy"] = {"value": f.value / f.columns["q_ly"] - 1, "how": f"{k} vs the same quarter a year earlier (p. {f.page})"}
            fy = ((last_fy or {}).get(k) or {}).get("columns", {}).get("ytd")
            if k in FLOWS and fy is not None and f.columns.get("ytd") is not None and f.columns.get("ytd_ly") is not None:
                m[f"{k}_ltm"] = {"value": fy + f.columns["ytd"] - f.columns["ytd_ly"],
                                 "how": f"last twelve months: {year - 1} (Q4 {year - 1} report) + YTD {year} - YTD {year - 1}"}
    rev, pers = figs.get("pl_revenue"), figs.get("pl_personnel")
    if rev and pers:
        qcol = f"Q{q} {year}"
        if rev.columns.get(qcol) and pers.columns.get(qcol):
            m["personnel_pct_q"] = {"value": pers.columns[qcol] / rev.columns[qcol], "how": f"personnel / revenue, Q{q} (p. {pers.page})"}
        ly = f"Q{q} {year - 1}"
        if rev.columns.get(ly) and pers.columns.get(ly):
            m["personnel_pct_q_ly"] = {"value": pers.columns[ly] / rev.columns[ly], "how": f"personnel / revenue, Q{q} {year - 1}"}
        lr, lp = ltm(rev.columns, year, q), ltm(pers.columns, year, q)
        if lr and lp:
            m["revenue_ltm"] = {"value": lr, "how": "revenue, last twelve months: FY + YTD - YTD a year earlier"}
            m["personnel_pct_ltm"] = {"value": lp / lr, "how": "personnel / revenue, last twelve months"}
    for k in ("ebitda_pre16", "ebit_pre16"):
        if m.get(f"{k}_ltm") and m.get("revenue_ltm"):
            m[f"{k}_margin_ltm"] = {"value": m[f"{k}_ltm"]["value"] / m["revenue_ltm"]["value"], "how": f"{k} / revenue, last twelve months"}
    return m


def last_fy_figures(code: str, year: int) -> dict | None:
    """The key figures of the Q4 report of ``year`` (its YTD column is the full year), if it has been read."""
    d = jg.home() / "reports" / code
    for p in d.glob("*.summary.json") if d.exists() else []:
        s = json.loads(p.read_text(encoding="utf-8"))
        if s.get("year") == year and s.get("quarter") == 4:
            return s["figures"]
    return None


# ------------------------------------------------------------------------------------------------ consensus (MFN)
def fetch_consensus(slug: str, year: int, q: int) -> dict | None:
    """The consensus a company publishes through Modular Finance before its report (MNOK unless stated)."""
    url = MFN.format(slug=slug, q=q, year=year)
    try:
        html = fd._get(url).decode("utf-8", "ignore")
    except Exception:  # noqa: BLE001 – no release is a normal answer
        return None
    text = re.sub(r"\s*\n\s*", "|", re.sub(r"<[^>]+>", "\n", re.sub(r"<script.*?</script>|<style.*?</style>", "", html, flags=re.S)))
    if f"Q{q} {year}E" not in text:
        return None
    m = re.search(r"\|(\d{4}-\d{2}-\d{2} \d{2}:\d{2}):\d{2}\|Estimates compiled", text)
    n = re.search(r"Number of contributors: (\d+)", text)
    seg = text[text.find(f"Q{q} {year}E"):]
    seg = seg[:seg.find("More consensus")] if "More consensus" in seg else seg[:800]
    parts = [p for p in seg.split("|")[1:] if p.strip()]
    kpis = {}
    for name, val in zip(parts[0::2], parts[1::2], strict=False):
        try:
            kpis[name.strip()] = float(val.replace(",", ""))
        except ValueError:
            continue
    names = {"Total revenue": "revenue", "EBITDA": "ebitda", "EBITDA before impact of IFRS 16": "ebitda_pre16",
             "Profit/loss for the period": "profit", "Earnings per share (NOK)": "eps", "Members (T)": "members_k"}
    return {"values": {names[k]: v for k, v in kpis.items() if k in names}, "published": m.group(1) if m else None,
            "contributors": int(n.group(1)) if n else None, "url": url, "source": "Modular Finance on behalf of the company, via MFN"}


# ------------------------------------------------------------------------------------------------ case assumptions
def case_values(model: dict | None, cfg: dict, year: int) -> dict:
    """The case model's numbers for the report's year (EST): read from the model sheet by the rows in its map."""
    if not model or not cfg:
        return {}
    import openpyxl

    xlsx, cmap = Path(model["xlsx"]).expanduser(), Path(model["map"]).expanduser()
    if not xlsx.exists() or not cmap.exists():
        return {}
    rows = json.loads(cmap.read_text(encoding="utf-8"))["rows"]
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    sheet = cfg.get("sheet", "Model")
    ws = wb[sheet]
    hdr = [c.value for c in next(ws.iter_rows(min_row=cfg.get("year_row", 3), max_row=cfg.get("year_row", 3)))]
    if year not in hdr:
        return {}
    col = hdr.index(year)
    out = {}
    for key, row_key in cfg.items():
        if key in ("sheet", "year_row"):
            continue
        r = rows.get(f"{sheet}|{row_key}")
        if r is None:
            continue
        v = next(ws.iter_rows(min_row=r, max_row=r, values_only=True))[col]
        if isinstance(v, (int, float)):
            out[key] = float(v)
    if "personnel" in out and "revenue" in out and out["revenue"]:
        out["personnel_pct"] = abs(out["personnel"]) / out["revenue"]
    out["_file"] = xlsx.name
    return out


# ------------------------------------------------------------------------------------------------ kill rules
def kill_status(rule: dict, now: float | None, before: float | None) -> str:
    if now is None:
        return "not measurable from the report"
    hit = now > rule["value"] if rule["op"] == ">" else now < rule["value"]
    if hit:
        return "triggered"
    if before is None:
        return "intact"
    d_now, d_before = abs(now - rule["value"]), abs(before - rule["value"])
    if abs(d_now - d_before) < 1e-9:
        return "intact, unchanged"
    return "intact, nearer" if d_now < d_before else "intact, further away"


# ------------------------------------------------------------------------------------------------ the pipeline
def find_reports(code: str, since: date, until: date | None = None) -> list[dict]:
    until = until or datetime.now(UTC).date()
    url = f"{fd.NEWSWEB}/list?issuer={code}&fromDate={since:%Y-%m-%d}&toDate={until:%Y-%m-%d}"
    msgs = json.loads(fd._get(url, method="POST"))["data"]["messages"]
    out = []
    for m in msgs:
        cats = {c.get("id") for c in m.get("category") or []}
        if m.get("numbAttachments") and period_of(m["title"]) and (cats & REPORT_CATEGORIES or re.search(r"(?i)report|rapport", m["title"])):
            if not re.search(r"(?i)invitation|presentation of", m["title"]):
                out.append(m)
    return sorted(out, key=lambda m: m["publishedTime"])


def download(code: str, message_id: int) -> tuple[Path, dict]:
    d = report_dir(code)
    meta_p = d / f"{message_id}.json"
    if not meta_p.exists():
        meta_p.write_bytes(fd._get(f"{fd.NEWSWEB}/message?messageId={message_id}"))
    msg = json.loads(meta_p.read_text(encoding="utf-8"))["data"]["message"]
    pdf = d / f"{message_id}.pdf"
    if not pdf.exists():
        atts = msg.get("attachments") or []
        pick = ([a for a in atts if re.search(r"(?i)report|rapport", a["name"]) and not re.search(r"(?i)presentation", a["name"])]
                or [a for a in atts if a["name"].lower().endswith(".pdf")])
        if not pick:
            raise ValueError(f"no report PDF in message {message_id}")
        pdf.write_bytes(fd._get(f"{fd.NEWSWEB}/attachment?messageId={message_id}&attachmentId={pick[0]['id']}"))
    return pdf, msg


def analyse(code: str, message_id: int, spec: dict, model: dict | None, kill_texts: list[str], previous: dict | None = None,
            consensus: dict | None | bool = True) -> dict:
    pdf, msg = download(code, message_id)
    year, q = period_of(msg["title"])
    figs = extract(pdf, spec)
    metrics = derive(figs, year, q, last_fy_figures(code, year - 1))
    if consensus is True:
        consensus = fetch_consensus(spec["consensus"]["mfn"], year, q) if (spec.get("consensus") or {}).get("mfn") else None
    case = case_values(model, spec.get("case") or {}, year)
    prev_m = (previous or {}).get("metrics", {})
    rows = []
    for rule in spec.get("kill") or []:
        now = (metrics.get(rule["metric"]) or {}).get("value")
        before = (prev_m.get(rule["metric"]) or {}).get("value")
        text = kill_texts[rule["criterion"]] if isinstance(rule.get("criterion"), int) and rule["criterion"] < len(kill_texts) else rule.get("text", "")
        rows.append({"criterion": text, "metric": rule["metric"], "op": rule["op"], "threshold": rule["value"], "now": now,
                     "previous": before, "status": kill_status(rule, now, before), "note": rule.get("note", "")})
    covered = {r["criterion"] for r in rows}
    for t in kill_texts:
        if t not in covered:
            rows.append({"criterion": t, "metric": None, "status": "not machine-checkable: read the report", "now": None})
    return {"code": code, "message_id": message_id, "title": msg["title"], "published": msg["publishedTime"], "year": year, "quarter": q,
            "url": f"https://newsweb.oslobors.no/message/{message_id}", "pdf": pdf.name,
            "figures": {k: {"value": f.value, "page": f.page, "line": f.line, "columns": f.columns, "checked": f.checked} for k, f in figs.items()},
            "metrics": metrics, "previous": {"title": previous["title"], "metrics": prev_m} if previous else None,
            "consensus": consensus, "case": case, "kill": rows, "analysed_at": datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")}


def summary_path(code: str, message_id: int) -> Path:
    return report_dir(code) / f"{message_id}.summary.json"


def run(code: str, model: dict | None, kill_texts: list[str], days: int = 400, force: bool = False) -> list[dict]:
    """Analyse every report of the last ``days`` days not yet analysed (oldest first, so each has its predecessor)."""
    spec = specs().get(code)
    if not spec:
        return []
    out, previous = [], None
    for m in find_reports(code, datetime.now(UTC).date() - timedelta(days=days)):
        p = summary_path(code, m["messageId"])
        if p.exists() and not force:
            previous = json.loads(p.read_text(encoding="utf-8"))
            continue
        s = analyse(code, m["messageId"], spec, model, kill_texts, previous)
        p.write_text(json.dumps(s, ensure_ascii=False, indent=1), encoding="utf-8")
        out.append(s)
        previous = s
    return out


def latest(code: str) -> dict | None:
    d = jg.home() / "reports" / code
    if not d.exists():
        return None
    sums = [json.loads(p.read_text(encoding="utf-8")) for p in d.glob("*.summary.json")]
    return max(sums, key=lambda s: s["published"]) if sums else None
